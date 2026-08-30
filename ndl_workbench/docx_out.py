"""Render the translation markdown to .docx.

The original pipeline drove Word over COM. This does the same job with
python-docx so the packaged app has no dependency on a Word install, and so it
runs the same way on a machine that has never opened Office.

Frame headings become Heading 2, which makes Word's navigation pane a clickable
frame index - the single feature that makes a 100-page translation usable.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

_INLINE = re.compile(r"(\*\*.+?\*\*|(?<![\w*])\*[^*\r\n]+?\*(?![\w*])|`[^`]+?`)")


def _apply_fonts(run, latin: str, cjk: str) -> None:
    run.font.name = latin
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), cjk)


def _add_inline(paragraph, text: str, latin: str, cjk: str) -> None:
    """Write text into a paragraph, honouring **bold**, *italic* and `code`."""
    for part in _INLINE.split(text):
        if not part:
            continue
        bold = italic = mono = False
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            part, bold = part[2:-2], True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            part, mono = part[1:-1], True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            part, italic = part[1:-1], True
        run = paragraph.add_run(part)
        run.bold = bold
        run.italic = italic
        _apply_fonts(run, "Consolas" if mono else latin, cjk)


def _add_page_number_footer(section, latin: str, cjk: str) -> None:
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    _apply_fonts(run, latin, cjk)
    for el, attrs, text in (
        ("w:fldChar", {"w:fldCharType": "begin"}, None),
        ("w:instrText", {"xml:space": "preserve"}, " PAGE "),
        ("w:fldChar", {"w:fldCharType": "end"}, None),
    ):
        node = OxmlElement(el)
        for k, v in attrs.items():
            node.set(qn(k), v)
        if text is not None:
            node.text = text
        run._element.append(node)


def _split_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def markdown_to_docx(md_path: Path, docx_path: Path,
                     *, latin_font: str = "Georgia", cjk_font: str = "Yu Mincho") -> Path:
    """Convert the subset of markdown this pipeline emits."""
    md = md_path.read_text(encoding="utf-8")
    doc = Document()

    for style_name in ("Normal", "Heading 1", "Heading 2", "Heading 3"):
        try:
            style = doc.styles[style_name]
        except KeyError:
            continue
        style.font.name = latin_font
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        rfonts.set(qn("w:ascii"), latin_font)
        rfonts.set(qn("w:hAnsi"), latin_font)
        rfonts.set(qn("w:eastAsia"), cjk_font)
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    doc.styles["Heading 2"].paragraph_format.keep_with_next = True

    _add_page_number_footer(doc.sections[0], latin_font, cjk_font)

    lines = md.splitlines()
    i = 0
    para: list[str] = []

    def flush() -> None:
        nonlocal para
        if para:
            p = doc.add_paragraph()
            _add_inline(p, " ".join(para), latin_font, cjk_font)
            para = []

    while i < len(lines):
        line = lines[i].rstrip()

        if line.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                raw = lines[i].strip()
                if not re.fullmatch(r"\|[\s:|-]+\|", raw):
                    rows.append(_split_row(raw))
                i += 1
            if rows:
                width = max(len(r) for r in rows)
                table = doc.add_table(rows=0, cols=width)
                table.style = "Table Grid"
                for r in rows:
                    cells = table.add_row().cells
                    for j in range(width):
                        cell = cells[j]
                        cell.text = ""
                        _add_inline(cell.paragraphs[0], r[j] if j < len(r) else "",
                                    latin_font, cjk_font)
            continue

        m = re.match(r"^(#{1,3})\s+(.*)", line)
        if m:
            flush()
            level = len(m.group(1))
            h = doc.add_heading(level=level)
            _add_inline(h, m.group(2), latin_font, cjk_font)
            if level == 1:
                h.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        if re.fullmatch(r"(-{3,}|\*{3,})", line.strip()):
            flush()
            doc.add_paragraph().add_run().add_break()
            i += 1
            continue

        m = re.match(r"^\s*[-*]\s+(.*)", line)
        if m:
            flush()
            p = doc.add_paragraph(style="List Bullet")
            _add_inline(p, m.group(1), latin_font, cjk_font)
            i += 1
            continue

        m = re.match(r"^\s*(\d+)\.\s+(.*)", line)
        if m:
            flush()
            p = doc.add_paragraph(style="List Number")
            _add_inline(p, m.group(2), latin_font, cjk_font)
            i += 1
            continue

        if line.startswith(">"):
            flush()
            p = doc.add_paragraph(style="Intense Quote" if "Intense Quote" in
                                  [s.name for s in doc.styles] else "Normal")
            _add_inline(p, line.lstrip("> ").rstrip(), latin_font, cjk_font)
            i += 1
            continue

        if not line.strip():
            flush()
            i += 1
            continue

        para.append(line.strip())
        i += 1

    flush()
    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(docx_path))
    return docx_path
